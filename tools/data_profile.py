#!/usr/bin/env python3
"""为 CSV、JSON 或 XLSX 表格数据生成紧凑的 JSON 概览。"""
from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from common import write_json


def parse_value(raw: Any) -> Any:
    if raw is None or raw == "":
        return None
    if isinstance(raw, (int, float, bool)):
        return raw
    text = str(raw).strip()
    try:
        return int(text)
    except ValueError:
        try:
            return float(text)
        except ValueError:
            return text


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [{k: parse_value(v) for k, v in row.items()} for row in csv.DictReader(handle)]
    if path.suffix.lower() == ".json":
        with path.open("r", encoding="utf-8") as handle:
            value = json.load(handle)
        if not isinstance(value, list) or any(not isinstance(x, dict) for x in value):
            raise ValueError("JSON 输入必须是对象数组")
        return [{str(k): parse_value(v) for k, v in row.items()} for row in value]
    raise ValueError("read_rows 仅用于 .csv 和 .json 输入")


def read_xlsx_sheets(path: Path) -> list[tuple[str, list[dict[str, Any]]]]:
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ValueError("读取 XLSX 需要可选依赖 openpyxl；请安装后重试") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: list[tuple[str, list[dict[str, Any]]]] = []
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if header_row is None:
                sheets.append((sheet.title, []))
                continue
            headers: list[str] = []
            seen: Counter[str] = Counter()
            for index, value in enumerate(header_row, start=1):
                base = str(value).strip() if value is not None and str(value).strip() else f"column_{get_column_letter(index)}"
                seen[base] += 1
                headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
            rows = [
                {headers[index]: parse_value(value) for index, value in enumerate(row)}
                for row in iterator
                if any(value is not None and value != "" for value in row)
            ]
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def iso_date(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    candidate = value.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(candidate).isoformat()
    except ValueError:
        return None


def profile_rows(rows: list[dict[str, Any]], source: str) -> dict[str, Any]:
    columns = sorted({key for row in rows for key in row})
    signatures = [json.dumps(row, ensure_ascii=False, sort_keys=True, default=str) for row in rows]
    duplicate_rows = sum(count - 1 for count in Counter(signatures).values() if count > 1)
    summaries: dict[str, Any] = {}
    for name in columns:
        values = [row.get(name) for row in rows]
        present = [value for value in values if value is not None]
        numeric = [float(value) for value in present if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))]
        types = Counter(type(value).__name__ for value in present)
        item: dict[str, Any] = {
            "dtype": types.most_common(1)[0][0] if types else "null",
            "missing": len(values) - len(present),
            "unique_count": len({json.dumps(value, ensure_ascii=False, sort_keys=True, default=str) for value in present}),
        }
        if numeric and len(numeric) == len(present):
            mean = statistics.fmean(numeric)
            std = statistics.stdev(numeric) if len(numeric) > 1 else 0.0
            item["numeric"] = {"min": min(numeric), "max": max(numeric), "mean": mean, "std": std}
            item["obvious_outlier_count"] = sum(abs(x - mean) > 3 * std for x in numeric) if std else 0
        else:
            counts = Counter(str(value) for value in present)
            item["categorical_summary"] = [{"value": key, "count": count} for key, count in counts.most_common(10)]
            dates = [iso_date(value) for value in present]
            parsed = [value for value in dates if value is not None]
            if parsed and len(parsed) == len(present):
                item["time_range"] = {"min": min(parsed), "max": max(parsed)}
        nonempty = len(present)
        item["possible_id"] = bool(nonempty and item["unique_count"] == nonempty and ("id" in name.lower() or nonempty == len(rows)))
        summaries[name] = item
    return {"source": source, "shape": [len(rows), len(columns)], "columns": summaries, "suspicious_duplicate_rows": duplicate_rows}


def profile(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".json"}:
        return profile_rows(read_rows(path), str(path))
    if suffix == ".xlsx":
        sheets = [profile_rows(rows, f"{path}#{name}") | {"sheet": name} for name, rows in read_xlsx_sheets(path)]
        return {"source": str(path), "format": "xlsx", "sheet_count": len(sheets), "sheets": sheets}
    raise ValueError("仅支持 .csv、.json 和 .xlsx 输入")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="CSV、XLSX 或对象数组形式的 JSON 文件")
    parser.add_argument("--output", type=Path, default=Path("workspace/data/data_summary.json"))
    args = parser.parse_args()
    write_json(args.output, profile(args.input))
    print(f"已写入 {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
