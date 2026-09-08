#!/usr/bin/env python3
"""2025 C题路线A/B的可复现实验运行器（仅依赖 numpy 与 openpyxl）。"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def load_config(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    required = {"data_path", "model_route", "folds", "threshold", "risk_fail", "risk_delay", "min_group_size", "seed"}
    missing = required - value.keys()
    if missing:
        raise ValueError(f"配置缺少字段：{sorted(missing)}")
    if value["model_route"] not in {"A", "B"}:
        raise ValueError("model_route 必须为 A 或 B")
    return value


def read_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    raw_headers = next(rows)
    headers = [str(x).strip() if x is not None else f"__blank_{i}" for i, x in enumerate(raw_headers)]
    return [dict(zip(headers, row, strict=True)) for row in rows]


def num(value: Any) -> float:
    try:
        x = float(value)
        return x if math.isfinite(x) else float("nan")
    except (TypeError, ValueError):
        return float("nan")


def week_num(value: Any) -> float:
    text = str(value).strip().lower()
    if "w" not in text:
        return num(value)
    weeks, rest = text.split("w", 1)
    days = rest.replace("+", "").strip() or "0"
    return float(weeks) + float(days) / 7.0


def sigmoid(z: np.ndarray) -> np.ndarray:
    z = np.clip(z, -35.0, 35.0)
    return 1.0 / (1.0 + np.exp(-z))


def standardize_fit(x: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    mean = np.nanmean(x, axis=0)
    scale = np.nanstd(x, axis=0)
    scale[scale < 1e-10] = 1.0
    return mean, scale


def standardize_apply(x: np.ndarray, mean: np.ndarray, scale: np.ndarray) -> np.ndarray:
    clean = np.where(np.isfinite(x), x, mean)
    return (clean - mean) / scale


def spline_basis(x: np.ndarray, knots: list[float]) -> np.ndarray:
    cols = [x, x * x]
    cols.extend(np.maximum(x - knot, 0.0) ** 3 for knot in knots)
    return np.column_stack(cols)


def design_concentration(rows: list[dict[str, Any]], route: str) -> tuple[np.ndarray, np.ndarray, np.ndarray, list[str]]:
    week = np.array([week_num(r["检测孕周"]) for r in rows])
    bmi = np.array([num(r["孕妇BMI"]) for r in rows])
    y = np.array([num(r["Y染色体浓度"]) for r in rows])
    groups = np.array([str(r["孕妇代码"]) for r in rows])
    if route == "A":
        x = np.column_stack([week, bmi, week * bmi])
        names = ["孕周", "BMI", "孕周×BMI"]
    else:
        wk = [float(v) for v in np.quantile(week, [0.25, 0.5, 0.75])]
        bk = [float(v) for v in np.quantile(bmi, [0.33, 0.67])]
        sw, sb = spline_basis(week, wk), spline_basis(bmi, bk)
        x = np.column_stack([sw, sb, week * bmi])
        names = ["孕周", "孕周²"] + [f"孕周截断三次项@{v:.2f}" for v in wk]
        names += ["BMI", "BMI²"] + [f"BMI截断三次项@{v:.2f}" for v in bk] + ["孕周×BMI"]
    return x, y, groups, names


def ridge_random_intercept(x: np.ndarray, y: np.ndarray, groups: np.ndarray, penalty: float) -> tuple[np.ndarray, dict[str, float]]:
    mean, scale = standardize_fit(x)
    z = standardize_apply(x, mean, scale)
    ids = sorted(set(groups.tolist()))
    pos = {g: i for i, g in enumerate(ids)}
    d = np.zeros((len(y), len(ids)))
    d[np.arange(len(y)), [pos[g] for g in groups]] = 1.0
    a = np.column_stack([np.ones(len(y)), z, d])
    reg = np.zeros(a.shape[1])
    reg[1 : 1 + z.shape[1]] = penalty * 0.05
    reg[1 + z.shape[1] :] = penalty
    beta = np.linalg.solve(a.T @ a + np.diag(reg) + np.eye(a.shape[1]) * 1e-9, a.T @ y)
    return beta[: 1 + z.shape[1]], {"mean": mean.tolist(), "scale": scale.tolist()}


def predict_fixed(x: np.ndarray, beta: np.ndarray, scaler: dict[str, list[float]]) -> np.ndarray:
    z = standardize_apply(x, np.array(scaler["mean"]), np.array(scaler["scale"]))
    return np.column_stack([np.ones(len(z)), z]) @ beta


def group_folds(groups: np.ndarray, folds: int, seed: int) -> list[np.ndarray]:
    unique = np.array(sorted(set(groups.tolist())))
    rng = np.random.default_rng(seed)
    rng.shuffle(unique)
    buckets = np.array_split(unique, folds)
    return [np.isin(groups, bucket) for bucket in buckets]


def regression_cv(x: np.ndarray, y: np.ndarray, groups: np.ndarray, route: str, folds: int, seed: int) -> dict[str, float]:
    predictions = np.full(len(y), np.nan)
    for valid in group_folds(groups, folds, seed):
        beta, scaler = ridge_random_intercept(x[~valid], y[~valid], groups[~valid], 1.0 if route == "B" else 3.0)
        predictions[valid] = predict_fixed(x[valid], beta, scaler)
    residual = y - predictions
    return {
        "grouped_cv_rmse": float(np.sqrt(np.mean(residual**2))),
        "grouped_cv_mae": float(np.mean(np.abs(residual))),
        "grouped_cv_r2": float(1.0 - np.sum(residual**2) / np.sum((y - np.mean(y)) ** 2)),
    }


def woman_records(rows: list[dict[str, Any]], threshold: float) -> list[dict[str, Any]]:
    by_id: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        by_id.setdefault(str(r["孕妇代码"]), []).append(r)
    output = []
    for woman, records in sorted(by_id.items()):
        records.sort(key=lambda r: week_num(r["检测孕周"]))
        times = np.array([week_num(r["检测孕周"]) for r in records])
        ys = np.array([num(r["Y染色体浓度"]) for r in records])
        hit = np.flatnonzero(ys >= threshold)
        if len(hit) == 0:
            censor, left, right = "right", float(times[-1]), None
        elif hit[0] == 0:
            censor, left, right = "left", 9.0, float(times[0])
        else:
            censor, left, right = "interval", float(times[hit[0] - 1]), float(times[hit[0]])
        base = records[0]
        output.append({
            "id": woman, "censor": censor, "left": left, "right": right,
            "bmi": num(base["孕妇BMI"]), "age": num(base["年龄"]), "height": num(base["身高"]),
            "weight": num(base["体重"]), "ivf": 0.0 if str(base["IVF妊娠"]) == "自然受孕" else 1.0,
            "gc": float(np.nanmedian([num(r["GC含量"]) for r in records])),
            "map": float(np.nanmedian([num(r["在参考基因组上比对的比例"]) for r in records])),
            "filter": float(np.nanmedian([num(r["被过滤掉读段数的比例"]) for r in records])),
        })
    return output


def hazard_raw(women: list[dict[str, Any]], weeks: np.ndarray) -> tuple[np.ndarray, list[str]]:
    blocks = []
    for p in women:
        t = weeks
        bmi = np.full(len(t), p["bmi"])
        blocks.append(np.column_stack([t, t * t, np.maximum(t - 14, 0) ** 3, np.maximum(t - 18, 0) ** 3,
                                       bmi, bmi * bmi, t * bmi, np.full(len(t), p["age"]),
                                       np.full(len(t), p["height"]), np.full(len(t), p["weight"]),
                                       np.full(len(t), p["ivf"]), np.full(len(t), p["gc"]),
                                       np.full(len(t), p["map"]), np.full(len(t), p["filter"])]))
    names = ["week", "week2", "week_spline14", "week_spline18", "bmi", "bmi2", "week_bmi", "age", "height", "weight", "ivf", "gc", "map", "filter"]
    return np.stack(blocks), names


def fit_interval_hazard(women: list[dict[str, Any]], weeks: np.ndarray, penalty: float, iterations: int = 1000) -> tuple[np.ndarray, dict[str, list[float]], list[float]]:
    raw, _ = hazard_raw(women, weeks)
    mean, scale = standardize_fit(raw.reshape(-1, raw.shape[-1]))
    x = standardize_apply(raw, mean, scale)
    x = np.concatenate([np.ones((*x.shape[:2], 1)), x], axis=2)
    beta = np.zeros(x.shape[2])
    m = np.zeros_like(beta)
    v = np.zeros_like(beta)
    history = []
    for step in range(1, iterations + 1):
        h = sigmoid(x @ beta)
        s = np.cumprod(1.0 - h, axis=1)
        grad = np.zeros_like(beta)
        loglik = 0.0
        for i, person in enumerate(women):
            li = int(np.searchsorted(weeks, person["left"], side="right") - 1)
            if person["censor"] == "right":
                li = max(li, 0)
                loglik += math.log(max(s[i, li], 1e-12))
                weights = np.where(np.arange(len(weeks)) <= li, -h[i], 0.0)
            else:
                ri = int(np.searchsorted(weeks, person["right"], side="left"))
                ri = min(max(ri, 0), len(weeks) - 1)
                if person["censor"] == "left":
                    delta = max(1.0 - s[i, ri], 1e-12)
                    loglik += math.log(delta)
                    weights = np.where(np.arange(len(weeks)) <= ri, s[i, ri] * h[i] / delta, 0.0)
                else:
                    sl = 1.0 if li < 0 else s[i, li]
                    delta = max(sl - s[i, ri], 1e-12)
                    loglik += math.log(delta)
                    idx = np.arange(len(weeks))
                    weights = (-sl * h[i] * (idx <= li) + s[i, ri] * h[i] * (idx <= ri)) / delta
            grad += x[i].T @ weights
        grad -= penalty * np.r_[0.0, beta[1:]]
        grad /= len(women)
        m = 0.9 * m + 0.1 * grad
        v = 0.999 * v + 0.001 * grad * grad
        update = 0.03 * (m / (1 - 0.9**step)) / (np.sqrt(v / (1 - 0.999**step)) + 1e-8)
        beta += update
        if step % 100 == 0:
            history.append(float(loglik - 0.5 * penalty * np.sum(beta[1:] ** 2)))
    if not np.all(np.isfinite(beta)):
        raise RuntimeError("区间删失风险模型未得到有限参数")
    return beta, {"mean": mean.tolist(), "scale": scale.tolist()}, history


def predict_cdf(women: list[dict[str, Any]], weeks: np.ndarray, beta: np.ndarray, scaler: dict[str, list[float]]) -> np.ndarray:
    raw, _ = hazard_raw(women, weeks)
    x = standardize_apply(raw, np.array(scaler["mean"]), np.array(scaler["scale"]))
    x = np.concatenate([np.ones((*x.shape[:2], 1)), x], axis=2)
    return 1.0 - np.cumprod(1.0 - sigmoid(x @ beta), axis=1)


def optimize_groups(women: list[dict[str, Any]], cdf: np.ndarray, weeks: np.ndarray, min_size: int, risk_fail: float, risk_delay: float) -> dict[str, Any]:
    bmi = np.array([p["bmi"] for p in women])
    candidates = sorted(set(float(x) for x in np.quantile(bmi, np.arange(0.15, 0.91, 0.1))))
    delay = np.maximum(weeks - 12.0, 0.0) / 15.0
    best: dict[str, Any] | None = None
    import itertools
    for groups in range(3, 6):
        for cuts in itertools.combinations(candidates, groups - 1):
            labels = np.digitize(bmi, cuts)
            sizes = [int(np.sum(labels == g)) for g in range(groups)]
            if min(sizes) < min_size:
                continue
            total, details = 0.0, []
            for g in range(groups):
                mask = labels == g
                curve = np.mean(cdf[mask], axis=0)
                risk = risk_fail * (1.0 - curve) + risk_delay * delay
                k = int(np.argmin(risk))
                total += float(risk[k]) * sizes[g]
                details.append({"group": g + 1, "bmi_lower": None if g == 0 else cuts[g - 1], "bmi_upper": None if g == groups - 1 else cuts[g], "n": sizes[g], "week": float(weeks[k]), "attainment_probability": float(curve[k]), "risk": float(risk[k])})
            score = total / len(women) + 0.01 * groups
            if best is None or score < best["objective"]:
                best = {"objective": score, "cuts": list(cuts), "groups": details}
    if best is None:
        raise RuntimeError("没有满足最小样本量约束的BMI有序分组")
    return best


def fixed_groups(women: list[dict[str, Any]], cdf: np.ndarray, weeks: np.ndarray, risk_fail: float, risk_delay: float) -> dict[str, Any]:
    bmi = np.array([p["bmi"] for p in women])
    cuts = [28.0, 32.0, 36.0, 40.0]
    labels = np.digitize(bmi, cuts)
    delay = np.maximum(weeks - 12.0, 0.0) / 15.0
    details, total = [], 0.0
    for g in range(5):
        mask = labels == g
        if not np.any(mask):
            continue
        curve = np.mean(cdf[mask], axis=0)
        risk = risk_fail * (1.0 - curve) + risk_delay * delay
        k = int(np.argmin(risk))
        total += float(risk[k]) * int(mask.sum())
        details.append({"group": g + 1, "bmi_lower": None if g == 0 else cuts[g - 1], "bmi_upper": None if g == 4 else cuts[g], "n": int(mask.sum()), "week": float(weeks[k]), "attainment_probability": float(curve[k]), "risk": float(risk[k])})
    return {"objective": total / len(women), "cuts": cuts, "groups": details}


def fit_logistic(x: np.ndarray, y: np.ndarray, penalty: float, class_weight: bool = True, iterations: int = 1200) -> np.ndarray:
    beta = np.zeros(x.shape[1])
    pos = max(float(np.sum(y == 1)), 1.0)
    neg = max(float(np.sum(y == 0)), 1.0)
    weights = np.where(y == 1, len(y) / (2 * pos), len(y) / (2 * neg)) if class_weight else np.ones(len(y))
    for step in range(1, iterations + 1):
        p = sigmoid(x @ beta)
        grad = x.T @ (weights * (p - y)) / len(y) + penalty * np.r_[0.0, beta[1:]] / len(y)
        beta -= 0.15 / math.sqrt(step / 40 + 1) * grad
    return beta


def average_precision(y: np.ndarray, p: np.ndarray) -> float | None:
    positives = int(np.sum(y))
    if positives == 0:
        return None
    order = np.argsort(-p)
    ys = y[order]
    precision = np.cumsum(ys) / np.arange(1, len(y) + 1)
    return float(np.sum(precision * ys) / positives)


def classification_metrics(y: np.ndarray, p: np.ndarray, threshold: float) -> dict[str, float | None]:
    pred = p >= threshold
    tp, tn = int(np.sum(pred & (y == 1))), int(np.sum(~pred & (y == 0)))
    fp, fn = int(np.sum(pred & (y == 0))), int(np.sum(~pred & (y == 1)))
    return {"sensitivity": tp / (tp + fn) if tp + fn else None, "specificity": tn / (tn + fp) if tn + fp else None,
            "balanced_accuracy": 0.5 * (tp / (tp + fn) + tn / (tn + fp)) if tp + fn and tn + fp else None,
            "pr_auc": average_precision(y, p), "brier": float(np.mean((p - y) ** 2)), "tp": tp, "tn": tn, "fp": fp, "fn": fn}


def female_matrix(rows: list[dict[str, Any]]) -> tuple[np.ndarray, dict[str, np.ndarray], list[str], np.ndarray]:
    features = ["13号染色体的Z值", "18号染色体的Z值", "21号染色体的Z值", "X染色体的Z值", "GC含量", "原始读段数", "在参考基因组上比对的比例", "重复读段的比例", "唯一比对的读段数", "被过滤掉读段数的比例", "孕妇BMI", "年龄"]
    x = np.array([[num(r[f]) for f in features] for r in rows])
    labels = {tag: np.array([1.0 if tag in str(r["染色体的非整倍体"] or "") else 0.0 for r in rows]) for tag in ["T13", "T18", "T21"]}
    return x, labels, features, np.array([str(r["孕妇代码"]) for r in rows])


def q4_cv(rows: list[dict[str, Any]], folds: int, seed: int, penalty: float) -> dict[str, Any]:
    x, labels, features, ids = female_matrix(rows)
    fold_masks = group_folds(ids, folds, seed)
    output: dict[str, Any] = {"features": features, "n_records": len(ids), "n_women": len(set(ids.tolist())), "labels": {}}
    for tag, y in labels.items():
        pred = np.full(len(y), np.nan)
        chosen_thresholds = []
        for valid in fold_masks:
            mean, scale = standardize_fit(x[~valid])
            train = np.column_stack([np.ones(np.sum(~valid)), standardize_apply(x[~valid], mean, scale)])
            test = np.column_stack([np.ones(np.sum(valid)), standardize_apply(x[valid], mean, scale)])
            beta = fit_logistic(train, y[~valid], penalty)
            train_prob = sigmoid(train @ beta)
            thresholds = np.linspace(0.1, 0.9, 81)
            scores = [classification_metrics(y[~valid], train_prob, t)["balanced_accuracy"] or -1 for t in thresholds]
            chosen = float(thresholds[int(np.argmax(scores))])
            chosen_thresholds.append(chosen)
            pred[valid] = sigmoid(test @ beta)
        fold_predictions = np.zeros(len(y), dtype=bool)
        for valid, threshold in zip(fold_masks, chosen_thresholds, strict=True):
            fold_predictions[valid] = pred[valid] >= threshold
        binary_score = fold_predictions.astype(float)
        confusion = classification_metrics(y, binary_score, 0.5)
        probability_metrics = {"pr_auc": average_precision(y, pred), "brier": float(np.mean((pred - y) ** 2))}
        output["labels"][tag] = {"prevalence": float(np.mean(y)), "fold_thresholds": chosen_thresholds, **confusion, **probability_metrics}
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    cfg = load_config(args.config)
    data_path = Path(cfg["data_path"])
    male = read_sheet(data_path, "男胎检测数据")
    female = read_sheet(data_path, "女胎检测数据")
    route = cfg["model_route"]
    x, y, groups, names = design_concentration(male, route)
    q1_cv = regression_cv(x, y, groups, route, int(cfg["folds"]), int(cfg["seed"]))
    beta, scaler = ridge_random_intercept(x, y, groups, 1.0 if route == "B" else 3.0)
    women = woman_records(male, float(cfg["threshold"]))
    weeks = np.arange(10.0, 25.01, 0.5)
    hazard_beta, hazard_scaler, likelihood_history = fit_interval_hazard(women, weeks, 0.5 if route == "B" else 2.0)
    cdf = predict_cdf(women, weeks, hazard_beta, hazard_scaler)
    if route == "B":
        grouping = optimize_groups(women, cdf, weeks, int(cfg["min_group_size"]), float(cfg["risk_fail"]), float(cfg["risk_delay"]))
    else:
        grouping = fixed_groups(women, cdf, weeks, float(cfg["risk_fail"]), float(cfg["risk_delay"]))
    sensitivities = {}
    for risk_fail in [2.0, 3.0, 5.0]:
        choice = optimize_groups(women, cdf, weeks, int(cfg["min_group_size"]), risk_fail, float(cfg["risk_delay"])) if route == "B" else fixed_groups(women, cdf, weeks, risk_fail, float(cfg["risk_delay"]))
        sensitivities[f"risk_fail_{risk_fail:g}"] = choice
    q4 = q4_cv(female, int(cfg["folds"]), int(cfg["seed"]), 1.0 if route == "B" else 3.0)
    censor_counts = {kind: sum(p["censor"] == kind for p in women) for kind in ["left", "interval", "right"]}
    metrics = {"route": route, "q1": q1_cv, "q2_q3_objective": grouping["objective"], "q4": {k: {m: v for m, v in d.items() if m in {"prevalence", "sensitivity", "specificity", "balanced_accuracy", "pr_auc", "brier"}} for k, d in q4["labels"].items()}}
    results = {"route": route, "data": {"male_rows": len(male), "male_women": len(women), "female_rows": len(female), "female_women": q4["n_women"], "censor_counts": censor_counts},
               "q1": {"feature_names": names, "fixed_coefficients_standardized": beta.tolist(), "scaler": scaler, "cv": q1_cv},
               "q2_q3": {"week_grid": weeks.tolist(), "selected_grouping": grouping, "risk_sensitivity": sensitivities, "interval_loglik_history": likelihood_history}, "q4": q4,
               "limitations": ["附件样本偏向高BMI人群，分组不应外推至未覆盖人群", "离散风险模型使用0.5周网格", "Q4保留全部检测记录但严格按孕妇代码分折", "AE未作为预测特征"]}
    write_json(args.output / "metrics.json", metrics)
    write_json(args.output / "results.json", results)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
